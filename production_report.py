"""Formatted production report workbook (openpyxl).

The API gathers plain rows; this module only lays them out. Every figure is
written as a value, not a formula: downloaded workbooks open in Excel's
Protected View, which shows cached results only, so formulas would read as
blanks until the user clicks "Enable Editing".
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import date, datetime, timedelta
from statistics import median

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ── Palette ──────────────────────────────────────────────────────────────────
NAVY = "1F2A44"
GOLD = "C9A227"
INK = "1F2937"
MUTED = "6B7280"
LINE = "D1D5DB"
ZEBRA = "F8FAFC"
TOTAL_BG = "E5E7EB"
AM_BG, AM_FG = "FEF3C7", "92400E"
PM_BG, PM_FG = "EDE9FE", "5B21B6"
FLAG_BG, FLAG_FG = "FEE2E2", "991B1B"
OK_FG = "166534"

# Weights: kg below 1,000 kg, tons from 1,000 kg (matches the app's units rule).
KG = '[>=1000]#,##0.00," t";[=0]"–";#,##0.0" kg"'
KG_SIGNED = '#,##0.0" kg";[Red]-#,##0.0" kg";"–"'
INT = '#,##0;-#,##0;"–"'
COUNT = '0'
PCT = '0.0%'
HOURS = '0.0" h"'
DATE = 'dd mmm yyyy'
DATETIME = 'dd mmm yyyy hh:mm'
TIME = 'hh:mm'

THIN = Side(style="thin", color=LINE)
MEDIUM = Side(style="medium", color=NAVY)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TOP_SIZES = 15

FLOOR_LABEL = {"GROUND_FLOOR": "Ground Floor", "FIRST_FLOOR": "First Floor"}
LOCATION_LABEL = {"WAREHOUSE": "Warehouse", "GROUND_FLOOR": "Ground Floor",
                  "FIRST_FLOOR": "First Floor", "MACHINE": "Machine"}


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _num(value) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _ratio(part, whole):
    return (part / whole) if whole else None


def _day(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value)[:10])


def _fmt_day(d: date) -> str:
    return d.strftime("%d %b %Y")


# ── Low-level writers ────────────────────────────────────────────────────────

def _title(ws, text, subtitle, width_cols):
    last = get_column_letter(width_cols)
    ws.merge_cells(f"A1:{last}1")
    ws["A1"] = text
    ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(f"A2:{last}2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, italic=True, color=MUTED)
    ws["A2"].alignment = Alignment(indent=1)
    for col in range(1, width_cols + 1):
        ws.cell(row=1, column=col).fill = _fill(NAVY)
        ws.cell(row=3, column=col).border = Border(top=Side(style="thick", color=GOLD))


def _section(ws, row, col, text, span):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
    cell = ws.cell(row=row, column=col, value=text.upper())
    cell.font = Font(size=11, bold=True, color=NAVY)
    for c in range(col, col + span):
        ws.cell(row=row, column=c).border = Border(bottom=MEDIUM)
    return row + 1


def _style_shift(cell, shift):
    if shift == "AM":
        cell.fill, cell.font = _fill(AM_BG), Font(bold=True, color=AM_FG)
    elif shift == "PM":
        cell.fill, cell.font = _fill(PM_BG), Font(bold=True, color=PM_FG)
    cell.alignment = Alignment(horizontal="center")


def _table(ws, row, col, headers, data, formats=None, totals=None, shift_col=None,
           flag_col=None, bar_col=None, wrap_cols=(), group_col=None):
    """Write a boxed table; returns (first_data_row, next_free_row)."""
    formats = formats or {}
    for i, header in enumerate(headers):
        cell = ws.cell(row=row, column=col + i, value=header)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = _fill(NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = 30
    first = row + 1
    previous_group = None
    for r_idx, values in enumerate(data):
        r = first + r_idx
        group = values[group_col] if group_col is not None else None
        new_group = group_col is not None and r_idx > 0 and group != previous_group
        previous_group = group
        for i, value in enumerate(values):
            cell = ws.cell(row=r, column=col + i, value=value)
            # Operator names, sizes, notes and reasons are typed by users;
            # never let one starting with "=" become a live formula.
            if isinstance(value, str) and cell.data_type == "f":
                cell.data_type = "s"
            cell.border = Border(left=THIN, right=THIN, bottom=THIN,
                                 top=MEDIUM if new_group else THIN)
            if r_idx % 2:
                cell.fill = _fill(ZEBRA)
            if i in formats:
                cell.number_format = formats[i]
            if i in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if shift_col is not None:
            _style_shift(ws.cell(row=r, column=col + shift_col), values[shift_col])
        if flag_col is not None and values[flag_col]:
            for i in range(len(values)):
                ws.cell(row=r, column=col + i).fill = _fill(FLAG_BG)
            flag = ws.cell(row=r, column=col + flag_col)
            flag.font = Font(bold=True, color=FLAG_FG)
    last = first + len(data) - 1
    if not data:
        cell = ws.cell(row=first, column=col, value="No records in this period")
        cell.font = Font(italic=True, color=MUTED)
        ws.merge_cells(start_row=first, start_column=col, end_row=first, end_column=col + len(headers) - 1)
        last = first
    if bar_col is not None and data:
        letter = get_column_letter(col + bar_col)
        ws.conditional_formatting.add(
            f"{letter}{first}:{letter}{last}",
            DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
    nxt = last + 1
    if totals is not None:
        for i, value in enumerate(totals):
            cell = ws.cell(row=nxt, column=col + i, value=value)
            cell.font = Font(bold=True, color=INK)
            cell.fill = _fill(TOTAL_BG)
            cell.border = Border(left=THIN, right=THIN, top=MEDIUM, bottom=MEDIUM)
            if i in formats and value is not None and not isinstance(value, str):
                cell.number_format = formats[i]
        nxt += 1
    return first, nxt + 1


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _print_setup(ws, landscape=True, repeat_row=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if repeat_row:
        ws.print_title_rows = f"{repeat_row}:{repeat_row}"
    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "&A"


def _kpi(ws, row, col, label, value, fmt, hint, accent=GOLD):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    ws.merge_cells(start_row=row + 2, start_column=col, end_row=row + 2, end_column=col + 1)
    head = ws.cell(row=row, column=col, value=label.upper())
    head.font = Font(size=8, bold=True, color=MUTED)
    val = ws.cell(row=row + 1, column=col, value=value)
    val.font = Font(size=18, bold=True, color=INK)
    if fmt:
        val.number_format = fmt
    val.alignment = Alignment(horizontal="left", indent=1)
    head.alignment = Alignment(indent=1)
    sub = ws.cell(row=row + 2, column=col, value=hint)
    sub.font = Font(size=8, color=MUTED)
    sub.alignment = Alignment(indent=1)
    for r in range(row, row + 3):
        for c in range(col, col + 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = _fill("FFFFFF")
            cell.border = Border(
                left=Side(style="thick", color=accent) if c == col else None,
                right=THIN if c == col + 1 else None,
                top=THIN if r == row else None, bottom=THIN if r == row + 2 else None)
    ws.row_dimensions[row + 1].height = 26


# ── Data shaping ─────────────────────────────────────────────────────────────

def _entry_checks(entries):
    """Scrutiny flags per production entry id."""
    flags = defaultdict(list)
    dupes = defaultdict(list)
    by_machine_size = defaultdict(list)
    for e in entries:
        dupes[(e["business_date"], e["shift"], e["machine_id"], e["size"],
               round(_num(e["gross_weight"]), 3), round(_num(e["tare_weight"]), 3))].append(e["id"])
        by_machine_size[(e["machine_id"], e["size"])].append(_num(e["net_weight"]))
    for ids in dupes.values():
        if len(ids) > 1:
            for i in ids:
                flags[i].append(f"Possible duplicate ({len(ids)} identical)")
    medians = {k: median(v) for k, v in by_machine_size.items() if len(v) >= 4}
    for e in entries:
        net = _num(e["net_weight"])
        gross = _num(e["gross_weight"])
        if net <= 0:
            flags[e["id"]].append("Zero net weight")
        mid = medians.get((e["machine_id"], e["size"]))
        if mid and net > 0 and (net < 0.5 * mid or net > 1.5 * mid):
            flags[e["id"]].append(f"Net {net / mid:.0%} of usual for this machine & size")
        if gross and _num(e["tare_weight"]) / gross > 0.25:
            flags[e["id"]].append("Tare over 25% of gross")
        consumed = e.get("machine_consumed_kg")
        if consumed is not None and net > _num(consumed) + 0.001:
            flags[e["id"]].append("Net exceeds material consumed")
    return flags


def _lag_hours(entry):
    keyed, stamped = entry.get("keyed_at"), entry.get("created_at")
    if not keyed or not stamped:
        return None
    return max((keyed - stamped).total_seconds() / 3600, 0)


# ── Sheets ───────────────────────────────────────────────────────────────────

def _summary(wb, ctx):
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    entries, days = ctx["entries"], ctx["days"]
    _title(ws, "VENKATESWARA POLYMERS — PRODUCTION REPORT", ctx["subtitle"], 12)

    net_total = sum(_num(e["net_weight"]) for e in entries)
    am = sum(_num(e["net_weight"]) for e in entries if e["shift"] == "AM")
    pm = net_total - am
    consumed = sum(_num(e["machine_consumed_kg"]) for e in entries if e["machine_consumed_kg"] is not None)
    consumed_net = sum(_num(e["net_weight"]) for e in entries if e["machine_consumed_kg"] is not None)
    wastage = sum(_num(w["weight"]) for w in ctx["wastage"])
    active_machines = len({e["machine_id"] for e in entries})
    prod_days = len({e["business_date"] for e in entries})
    flagged = sum(1 for e in entries if ctx["flags"].get(e["id"]))

    _kpi(ws, 5, 1, "Net output", net_total, KG, f"{len(entries)} entries over {prod_days} day(s)")
    _kpi(ws, 5, 3, "Average per entry", _ratio(net_total, len(entries)) or 0, KG, "net weight per roll / entry")
    _kpi(ws, 5, 5, "AM · Morning", am, KG, f"{_ratio(am, net_total) or 0:.0%} of output", accent="F59E0B")
    _kpi(ws, 5, 7, "PM · Night", pm, KG, f"{_ratio(pm, net_total) or 0:.0%} of output", accent="7C3AED")
    _kpi(ws, 5, 9, "Wastage", wastage, KG,
         f"{_ratio(wastage, net_total) or 0:.1%} of net output", accent="E11D48")
    _kpi(ws, 5, 11, "Entries flagged", flagged, COUNT, "see Checks column", accent="E11D48" if flagged else "16A34A")

    _kpi(ws, 9, 1, "Material consumed", consumed, KG, "booked against machines")
    _kpi(ws, 9, 3, "Process yield", _ratio(consumed_net, consumed) or 0, PCT, "net ÷ consumed (where recorded)")
    _kpi(ws, 9, 5, "Process loss", consumed - consumed_net, KG, "consumed − net output")
    _kpi(ws, 9, 7, "Machines active", active_machines, COUNT, f"of {len(ctx['machines'])} machines")
    _kpi(ws, 9, 9, "Reversed entries", len(ctx["reversals"]), COUNT, "deleted after entry — see Corrections")
    _kpi(ws, 9, 11, "Production batches", len(ctx["batches"]), COUNT, "Ground Floor batches opened")

    row = 13
    # Daily output
    row = _section(ws, row, 1, "Daily output (by business date)", 12)
    wastage_by_day = defaultdict(float)
    for w in ctx["wastage"]:
        wastage_by_day[_day(w["date"])] += _num(w["weight"])
    data, tot = [], [0, 0, 0, 0, 0]
    for d in days:
        de = [e for e in entries if e["business_date"] == d]
        a = sum(_num(e["net_weight"]) for e in de if e["shift"] == "AM")
        p = sum(_num(e["net_weight"]) for e in de if e["shift"] == "PM")
        n_am = sum(1 for e in de if e["shift"] == "AM")
        n_pm = len(de) - n_am
        wst = wastage_by_day.get(d, 0.0)
        data.append([d, d.strftime("%a"), n_am, a, n_pm, p, len(de), a + p,
                     _ratio(a + p, net_total), wst, _ratio(wst, a + p)])
        for i, v in enumerate((n_am, a, n_pm, p, wst)):
            tot[i] += v
    _, row = _table(
        ws, row, 1,
        ["Date", "Day", "AM entries", "AM output", "PM entries", "PM output", "Total entries",
         "Total output", "Share", "Wastage", "Wastage %"],
        data,
        {0: DATE, 2: INT, 3: KG, 4: INT, 5: KG, 6: INT, 7: KG, 8: PCT, 9: KG, 10: PCT},
        totals=["TOTAL", "", tot[0], tot[1], tot[2], tot[3], tot[0] + tot[2], net_total,
                1 if net_total else None, tot[4], _ratio(tot[4], net_total)],
        bar_col=7)

    # Machine performance
    row = _section(ws, row, 1, "Machine performance", 12)
    data = []
    for m in ctx["machines"]:
        me = [e for e in entries if e["machine_id"] == m["id"]]
        if not me:
            continue
        a = sum(_num(e["net_weight"]) for e in me if e["shift"] == "AM")
        n = sum(_num(e["net_weight"]) for e in me)
        cons = sum(_num(e["machine_consumed_kg"]) for e in me if e["machine_consumed_kg"] is not None)
        cons_net = sum(_num(e["net_weight"]) for e in me if e["machine_consumed_kg"] is not None)
        active_days = len({e["business_date"] for e in me})
        data.append([m["name"], FLOOR_LABEL.get(m.get("floor"), m.get("floor") or "—"), len(me), a, n - a, n,
                     _ratio(n, net_total), _ratio(n, len(me)), n / active_days if active_days else None,
                     cons, _ratio(cons_net, cons)])
    data.sort(key=lambda r: -r[5])
    _, row = _table(
        ws, row, 1,
        ["Machine", "Floor", "Entries", "AM output", "PM output", "Net output", "Share",
         "Avg / entry", "Avg / active day", "Consumed", "Yield"],
        data, {2: INT, 3: KG, 4: KG, 5: KG, 6: PCT, 7: KG, 8: KG, 9: KG, 10: PCT},
        totals=["TOTAL", "", len(entries), am, pm, net_total, 1 if net_total else None,
                _ratio(net_total, len(entries)), None, consumed, _ratio(consumed_net, consumed)],
        bar_col=5)

    # Size mix + material side by side
    top = row
    row_l = _section(ws, top, 1, f"Size mix (top {TOP_SIZES})", 5)
    sizes = defaultdict(lambda: [0, 0.0])
    for e in entries:
        s = sizes[e["size"] or "—"]
        s[0] += 1
        s[1] += _num(e["net_weight"])
    data = sorted(([k, v[0], v[1], _ratio(v[1], net_total), _ratio(v[1], v[0])] for k, v in sizes.items()),
                  key=lambda r: -r[2])
    if len(data) > TOP_SIZES:
        rest = data[TOP_SIZES:]
        n, kg = sum(r[1] for r in rest), sum(r[2] for r in rest)
        data = data[:TOP_SIZES] + [[f"Other sizes ({len(rest)})", n, kg, _ratio(kg, net_total), _ratio(kg, n)]]
    _, row_l = _table(ws, row_l, 1, ["Size", "Entries", "Net output", "Share", "Avg / entry"], data,
                      {1: INT, 2: KG, 3: PCT, 4: KG}, bar_col=2)

    row_r = _section(ws, top, 7, "Material used in output", 6)
    mats = defaultdict(lambda: [0, 0.0, 0.0])
    for e in entries:
        s = mats[e["material_name"] or "—"]
        s[0] += 1
        s[1] += _num(e["net_weight"])
        s[2] += _num(e["machine_consumed_kg"])
    data = sorted(([k, v[0], v[1], v[2], _ratio(v[1], net_total)] for k, v in mats.items()),
                  key=lambda r: -r[2])
    _, row_r = _table(ws, row_r, 7, ["Material", "Entries", "Net output", "Consumed", "Share"], data,
                      {1: INT, 2: KG, 3: KG, 4: PCT}, bar_col=2)
    row = max(row_l, row_r)

    row = _section(ws, row, 1, "Operators", 12)
    ops = defaultdict(lambda: [0, 0.0, set(), set()])
    for e in entries:
        s = ops[(e["worker_name"] or "—").strip() or "—"]
        s[0] += 1
        s[1] += _num(e["net_weight"])
        s[2].add(e["machine_name"])
        s[3].add((e["business_date"], e["shift"]))
    data = sorted(([k, v[0], v[1], _ratio(v[1], net_total), len(v[3]), ", ".join(sorted(v[2]))]
                   for k, v in ops.items()), key=lambda r: -r[2])
    _, row = _table(ws, row, 1, ["Operator", "Entries", "Net output", "Share", "Shifts worked", "Machines"],
                    data, {1: INT, 2: KG, 3: PCT, 4: INT}, bar_col=2)

    _widths(ws, [22, 14, 12, 14, 12, 16, 20, 14, 13, 13, 12, 13])
    _print_setup(ws, landscape=False)
    ws.freeze_panes = "A4"


def _shift_matrix(wb, ctx):
    ws = wb.create_sheet("Shift Matrix")
    ws.sheet_view.showGridLines = False
    machines = [m for m in ctx["machines"] if any(e["machine_id"] == m["id"] for e in ctx["entries"])] \
        or ctx["machines"]
    width = 3 + len(machines)
    _title(ws, "OUTPUT BY SHIFT × MACHINE", "Net output per business date and shift. Darker green = higher output.", width)
    cell_net = defaultdict(float)
    for e in ctx["entries"]:
        cell_net[(e["business_date"], e["shift"], e["machine_id"])] += _num(e["net_weight"])
    data = []
    for d in ctx["days"]:
        for shift in ("AM", "PM"):
            vals = [cell_net.get((d, shift, m["id"]), 0.0) or None for m in machines]
            data.append([d, shift, *vals, sum(v or 0 for v in vals)])
        day_vals = [(cell_net.get((d, "AM", m["id"]), 0.0) + cell_net.get((d, "PM", m["id"]), 0.0)) or None
                    for m in machines]
        data.append([d, "DAY", *day_vals, sum(v or 0 for v in day_vals)])
    headers = ["Date", "Shift", *[m["name"] for m in machines], "Total"]
    fmts = {0: DATE, **{i: KG for i in range(2, width)}}
    totals_vals = [sum(cell_net.get((d, s, m["id"]), 0.0) for d in ctx["days"] for s in ("AM", "PM"))
                   for m in machines]
    first, _ = _table(ws, 4, 1, headers, data, fmts,
                      totals=["PERIOD TOTAL", "", *totals_vals, sum(totals_vals)], shift_col=1, group_col=0)
    # Day subtotal rows: bold grey; shift rows get the heat map.
    shift_rows = []
    for i, values in enumerate(data):
        r = first + i
        if values[1] == "DAY":
            for c in range(1, width + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(bold=True)
                cell.fill = _fill(TOTAL_BG)
            ws.cell(row=r, column=2).value = "Day total"
        else:
            shift_rows.append(r)
    for r in shift_rows:
        rng = f"C{r}:{get_column_letter(2 + len(machines))}{r}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF",
            end_type="max", end_color="63BE7B"))
    _widths(ws, [14, 11, *[14] * len(machines), 15])
    ws.freeze_panes = "C5"
    _print_setup(ws, repeat_row=4)


def _entries_sheet(wb, ctx):
    ws = wb.create_sheet("Production Entries")
    ws.sheet_view.showGridLines = False
    headers = ["#", "Business date", "Day", "Shift", "Entry time", "Keyed in at", "Lag",
               "Machine", "Floor", "Batch", "Material", "Size", "Operator",
               "Gross", "Tare", "Net", "Tare %", "Consumed", "Process loss",
               "Entered by", "Entry ID", "Checks"]
    _title(ws, "PRODUCTION ENTRIES — FULL DETAIL",
           "Every production entry in the period. 'Entry time' is when the work happened (as entered); "
           "'Keyed in at' is when it was typed into the system. Red rows need a look.", len(headers))
    data = []
    for i, e in enumerate(ctx["entries"], start=1):
        net, gross, tare = _num(e["net_weight"]), _num(e["gross_weight"]), _num(e["tare_weight"])
        consumed = e["machine_consumed_kg"]
        data.append([
            i, e["business_date"], e["business_date"].strftime("%a"), e["shift"], e["created_at"],
            e.get("keyed_at"), _lag_hours(e), e["machine_name"],
            FLOOR_LABEL.get(e.get("floor"), e.get("floor") or ""), e.get("batch_no") or "",
            e["material_name"] or "", e["size"] or "", e["worker_name"] or "",
            gross, tare, net, _ratio(tare, gross),
            _num(consumed) if consumed is not None else None,
            (_num(consumed) - net) if consumed is not None else None,
            e.get("entered_by") or "", e["id"], "; ".join(ctx["flags"].get(e["id"], [])),
        ])
    tot_net = sum(r[15] for r in data)
    tot_cons = sum(r[17] or 0 for r in data)
    first, _ = _table(
        ws, 4, 1, headers, data,
        {0: INT, 1: DATE, 4: DATETIME, 5: DATETIME, 6: HOURS, 13: KG, 14: KG, 15: KG, 16: PCT,
         17: KG, 18: KG_SIGNED, 20: INT},
        totals=["", f"{len(data)} entries", "", "", "", "", "", "", "", "", "", "", "",
                sum(r[13] for r in data), sum(r[14] for r in data), tot_net, None, tot_cons,
                tot_cons - sum(r[15] for r in data if r[17] is not None), "", "", ""],
        shift_col=3, flag_col=21, group_col=1, wrap_cols=(21,))
    if data:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{first + len(data) - 1}"
    _widths(ws, [5, 13, 6, 7, 16, 16, 8, 18, 12, 16, 18, 18, 18, 11, 10, 11, 8, 11, 12, 14, 9, 36])
    ws.freeze_panes = "E5"
    _print_setup(ws, repeat_row=4)


def _stock_sheet(wb, ctx):
    ws = wb.create_sheet("Stock Position")
    ws.sheet_view.showGridLines = False
    headers = ["Date", "Opening (plant)", "Received", "Consumed", "Returned", "Adjusted", "Closing (plant)",
               "Warehouse", "Ground Floor", "First Floor", "Machines", "Carry-over check"]
    _title(ws, "OPENING & CLOSING STOCK",
           "Plant totals per business day. Each day's closing is the next day's opening; "
           "the last column confirms the carry-over.", len(headers))
    data, prev_close = [], None
    for led in ctx["ledgers"]:
        t = led["totals"]
        plant = t["PLANT"]
        opening = plant.get("opening_kg", 0)
        check = "—" if prev_close is None else ("✓ matches" if abs(opening - prev_close) < 0.01 else
                                                 f"✗ off by {opening - prev_close:,.1f} kg")
        data.append([_day(led["business_date"]), opening, plant.get("received_kg", 0),
                     plant.get("consumed_kg", 0), plant.get("returned_kg", 0), plant.get("adjusted_kg", 0),
                     plant.get("closing_kg", 0), t["WAREHOUSE"].get("closing_kg", 0),
                     t["GROUND_FLOOR"].get("closing_kg", 0), t["FIRST_FLOOR"].get("closing_kg", 0),
                     t["MACHINES"].get("closing_kg", 0), check])
        prev_close = plant.get("closing_kg", 0)
    first, row = _table(ws, 4, 1, headers, data,
                        {0: DATE, 1: KG, 2: KG, 3: KG, 4: KG, 5: KG_SIGNED, 6: KG, 7: KG, 8: KG, 9: KG, 10: KG})
    for i, values in enumerate(data):
        cell = ws.cell(row=first + i, column=12)
        cell.font = Font(bold=True, color=OK_FG if values[11].startswith("✓") else
                         (FLAG_FG if values[11].startswith("✗") else MUTED))

    row = _section(ws, row, 1, "By material — whole period", len(headers))
    per = {}
    for idx, led in enumerate(ctx["ledgers"]):
        for m in led["materials"]:
            p = m["locations"]["PLANT"]
            slot = per.setdefault(m["material_name"] or m["key"], {
                "opening": None, "received": 0.0, "consumed": 0.0, "returned": 0.0, "adjusted": 0.0, "closing": 0.0})
            if slot["opening"] is None:
                slot["opening"] = p["opening_kg"] if idx == 0 else 0.0
            for k in ("received", "consumed", "returned", "adjusted"):
                slot[k] += p[f"{k}_kg"]
            slot["closing"] = p["closing_kg"]
    data = [[name, s["opening"] or 0, s["received"], s["consumed"], s["returned"], s["adjusted"], s["closing"]]
            for name, s in sorted(per.items(), key=lambda kv: kv[0].lower())]
    _table(ws, row, 1, ["Material", "Opening", "Received", "Consumed", "Returned", "Adjusted", "Closing"], data,
           {1: KG, 2: KG, 3: KG, 4: KG, 5: KG_SIGNED, 6: KG},
           totals=["TOTAL", *[sum(r[i] for r in data) for i in range(1, 7)]], bar_col=3)
    _widths(ws, [22, 15, 13, 13, 12, 13, 15, 13, 13, 13, 13, 20])
    ws.freeze_panes = "B5"
    _print_setup(ws, repeat_row=4)


def _simple_sheet(wb, name, title, subtitle, headers, data, formats, widths, totals=None,
                  shift_col=None, wrap_cols=(), bar_col=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    _title(ws, title, subtitle, len(headers))
    first, _ = _table(ws, 4, 1, headers, data, formats, totals=totals, shift_col=shift_col,
                      wrap_cols=wrap_cols, bar_col=bar_col)
    if data:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{first + len(data) - 1}"
    _widths(ws, widths)
    ws.freeze_panes = "A5"
    _print_setup(ws, repeat_row=4)
    return ws


def _notes(wb, ctx):
    ws = wb.create_sheet("Notes")
    ws.sheet_view.showGridLines = False
    _title(ws, "HOW TO READ THIS REPORT", ctx["subtitle"], 2)
    notes = [
        ("Business date", "The day a shift belongs to. The night (PM) shift runs past midnight but stays on the date it started."),
        ("AM / PM", f"AM · Morning runs {ctx['am_start']}–{ctx['pm_start']}; PM · Night runs {ctx['pm_start']} to {ctx['am_start']} the next morning."),
        ("Entry time", "When the work happened, as entered by the operator (backdated entries default to 09:00 for AM and 21:00 for PM unless a time was typed)."),
        ("Keyed in at / Lag", "When the entry was actually typed into the system, and how many hours after the entry time. Next-day entry of both shifts normally shows 12–30 h."),
        ("Weights", "Shown in kg; values of 1,000 kg and more are shown in tons (t). Cells hold the exact kg value — widen or reformat to see raw kg."),
        ("Consumed / Process loss", "Material booked off the machine for the entry, and consumed minus net output. Blank where the entry did not record consumption."),
        ("Yield", "Net output ÷ material consumed, only over entries that recorded consumption."),
        ("Checks", "Possible duplicate = same date, shift, machine, size, gross and tare. Unusual net = under 50% or over 150% of the median for that machine & size (needs 4+ entries). Tare over 25% of gross. Net exceeds consumed."),
        ("Corrections", "Production entries that were deleted after being saved. The original stays in the audit trail."),
        ("Carry-over check", "Confirms each day's opening equals the previous day's closing."),
    ]
    _, _ = _table(ws, 4, 1, ["Term", "Meaning"], [list(n) for n in notes], wrap_cols=(1,))
    _widths(ws, [24, 110])
    for r in range(5, 5 + len(notes)):
        ws.row_dimensions[r].height = 32
        ws.cell(row=r, column=1).font = Font(bold=True, color=NAVY)
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")


def build_production_workbook(*, start: date, end: date, generated_by: str, generated_at: datetime,
                              entries, machines, batches, machine_issues, wastage, reversals, ledgers,
                              am_start: str, pm_start: str) -> bytes:
    for e in entries:
        e["business_date"] = _day(e["business_date"])
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    n_days = len(days)
    ctx = {
        "entries": entries, "machines": machines, "batches": batches, "machine_issues": machine_issues,
        "wastage": wastage, "reversals": reversals, "ledgers": ledgers, "days": days,
        "flags": _entry_checks(entries), "am_start": am_start, "pm_start": pm_start,
        "subtitle": (f"Period {_fmt_day(start)} – {_fmt_day(end)} ({n_days} day{'s' if n_days != 1 else ''})"
                     f"  ·  Generated {generated_at:%d %b %Y %H:%M} IST by {generated_by}"),
    }
    wb = Workbook()
    _summary(wb, ctx)
    _shift_matrix(wb, ctx)
    _entries_sheet(wb, ctx)
    _stock_sheet(wb, ctx)

    out_by_batch = defaultdict(lambda: [0, 0.0])
    for e in entries:
        if e.get("batch_id"):
            out_by_batch[e["batch_id"]][0] += 1
            out_by_batch[e["batch_id"]][1] += _num(e["net_weight"])
    _simple_sheet(
        wb, "Batches", "PRODUCTION BATCHES (GROUND FLOOR)", ctx["subtitle"],
        ["Batch no", "Business date", "Shift", "Status", "Entries", "Net output", "Note", "Opened by", "Opened at"],
        [[b["batch_no"], _day(b["business_date"]), b["shift"], b["status"], out_by_batch[b["id"]][0],
          out_by_batch[b["id"]][1], b.get("note") or "", b.get("created_by_name") or "", b["created_at"]]
         for b in batches],
        {1: DATE, 4: INT, 5: KG, 8: DATETIME}, [20, 14, 8, 11, 9, 13, 30, 16, 17],
        shift_col=2, wrap_cols=(6,))

    _simple_sheet(
        wb, "Machine Issues", "MATERIAL ISSUED TO MACHINES", ctx["subtitle"],
        ["Transfer no", "Business date", "Shift", "Entry time", "From", "To machine", "Material",
         "Quantity", "Status", "Note", "Entered by"],
        [[t["transfer_no"], _day(t["business_date"]), t["shift"], t["created_at"],
          LOCATION_LABEL.get(t["from_location"], t["from_location"]), t.get("machine_name") or "",
          t.get("material_name") or "", _num(t["quantity_kg"]), t["status"], t.get("note") or "",
          t.get("created_by_name") or ""] for t in machine_issues],
        {1: DATE, 3: DATETIME, 7: KG}, [22, 14, 8, 17, 14, 18, 18, 12, 11, 28, 16],
        totals=["TOTAL", "", "", "", "", "", "",
                sum(_num(t["quantity_kg"]) for t in machine_issues if t["status"] == "ACTIVE"), "active only", "", ""],
        shift_col=2, wrap_cols=(9,), bar_col=7)

    output_by_day = defaultdict(float)
    for e in entries:
        output_by_day[e["business_date"]] += _num(e["net_weight"])
    wdata = [[_day(w["date"]), _num(w["weight"]), output_by_day.get(_day(w["date"]), 0.0),
              _ratio(_num(w["weight"]), output_by_day.get(_day(w["date"]), 0.0)), w.get("created_by_name") or ""]
             for w in wastage]
    _simple_sheet(
        wb, "Wastage", "WASTAGE", ctx["subtitle"],
        ["Date", "Wastage", "Net output that day", "Wastage %", "Entered by"], wdata,
        {0: DATE, 1: KG, 2: KG, 3: PCT}, [14, 13, 18, 12, 18],
        totals=["TOTAL", sum(r[1] for r in wdata), sum(output_by_day.values()),
                _ratio(sum(r[1] for r in wdata), sum(output_by_day.values())), ""], bar_col=1)

    _simple_sheet(
        wb, "Corrections", "DELETED / REVERSED PRODUCTION ENTRIES", ctx["subtitle"],
        ["Reversed at", "Business date", "Shift", "Entry ID", "Machine", "Material", "Quantity", "Reason", "Reversed by"],
        [[r["occurred_at"], _day(r["business_date"]), r["shift"], r["source_id"], r.get("machine_name") or "",
          r.get("material_name") or "", _num(r["quantity_kg"]), r.get("reason") or "", r.get("created_by_name") or ""]
         for r in reversals],
        {0: DATETIME, 1: DATE, 3: INT, 6: KG}, [17, 14, 8, 9, 16, 18, 12, 34, 16],
        shift_col=2, wrap_cols=(7,))
    _notes(wb, ctx)

    wb.properties.title = "Production Report"
    wb.properties.creator = "Venkateswara Polymers IMS"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
