import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from production_report import build_production_workbook


def _entry(i, day, shift, machine_id, gross, tare, size="12x18", consumed=None, hour=10):
    return {
        "id": i, "business_date": day, "shift": shift, "created_at": datetime(day.year, day.month, day.day, hour, 0),
        "keyed_at": datetime(day.year, day.month, day.day + 1, 9, 30), "machine_id": machine_id,
        "machine_name": f"Machine {machine_id}", "floor": "GROUND_FLOOR", "batch_id": None, "batch_no": None,
        "material_name": "LD Natural", "size": size, "worker_name": "Ravi",
        "gross_weight": Decimal(str(gross)), "tare_weight": Decimal(str(tare)),
        "net_weight": Decimal(str(gross - tare)), "machine_consumed_kg": consumed, "entered_by": "Clerk",
    }


def _ledger(day, opening, closing):
    loc = {"opening_kg": 0, "closing_kg": 0, "received_kg": 0, "consumed_kg": 0, "returned_kg": 0, "adjusted_kg": 0}
    plant = {**loc, "opening_kg": opening, "closing_kg": closing, "consumed_kg": opening - closing}
    return {
        "business_date": day.isoformat(),
        "totals": {"PLANT": plant, "WAREHOUSE": {**loc, "closing_kg": closing}, "GROUND_FLOOR": loc,
                   "FIRST_FLOOR": loc, "MACHINES": loc},
        "materials": [{"key": "m1", "material_name": "LD Natural",
                       "locations": {"PLANT": plant}}],
    }


def _build(entries):
    d1, d2 = date(2026, 10, 13), date(2026, 10, 14)
    content = build_production_workbook(
        start=d1, end=d2, generated_by="Owner", generated_at=datetime(2026, 10, 15, 9, 0),
        entries=entries, machines=[{"id": 1, "name": "Machine 1", "floor": "GROUND_FLOOR"},
                                   {"id": 2, "name": "Machine 2", "floor": "GROUND_FLOOR"}],
        batches=[], machine_issues=[], wastage=[{"date": d1, "weight": Decimal("4"), "created_by_name": "Clerk"}],
        reversals=[], ledgers=[_ledger(d1, 1000.0, 900.0), _ledger(d2, 900.0, 820.0)],
        am_start="08:00", pm_start="20:00")
    return load_workbook(io.BytesIO(content))


def test_workbook_has_every_sheet_and_totals():
    d1, d2 = date(2026, 10, 13), date(2026, 10, 14)
    entries = [_entry(1, d1, "AM", 1, 30, 1), _entry(2, d1, "PM", 2, 25, 1, hour=22),
               _entry(3, d2, "AM", 1, 31, 1, consumed=Decimal("32"))]
    wb = _build(entries)
    assert wb.sheetnames == ["Summary", "Shift Matrix", "Production Entries", "Stock Position",
                             "Batches", "Machine Issues", "Wastage", "Corrections", "Notes"]
    summary = wb["Summary"]
    assert summary["A6"].value == 83  # net output KPI: 29 + 24 + 30
    rows = list(wb["Production Entries"].iter_rows(min_row=5, max_row=7, values_only=True))
    assert [r[15] for r in rows] == [29, 24, 30]
    assert rows[0][6] == 23.5  # keyed in 23.5 h after the work
    carry = [r[11] for r in wb["Stock Position"].iter_rows(min_row=5, max_row=6, values_only=True)]
    assert carry == ["—", "✓ matches"]


def test_duplicate_entries_are_flagged():
    d1 = date(2026, 10, 13)
    wb = _build([_entry(1, d1, "AM", 1, 30, 1), _entry(2, d1, "AM", 1, 30, 1)])
    checks = [r[21] for r in wb["Production Entries"].iter_rows(min_row=5, max_row=6, values_only=True)]
    assert all("Possible duplicate" in c for c in checks)


def test_empty_period_still_builds():
    wb = _build([])
    assert wb["Production Entries"]["A5"].value == "No records in this period"
